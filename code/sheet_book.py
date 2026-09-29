"""Import columns H–I from the bet sheet, and build a template whose D column resets itself."""

from __future__ import annotations

import json
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from auth_storage import APP_DIR, _restrict

SHEET_FILE = APP_DIR / "sheet.json"
HEADER_SKIP = {"尾数", "生肖", "下注", "尾数/生肖/40", "数字"}
TAILS = [str(digit) for digit in range(10)]
ZODIACS = ["鼠", "牛", "虎", "兔", "龙", "蛇", "马", "羊", "猴", "鸡", "狗", "猪"]


@dataclass
class SheetTargets:
    tail: str | None
    zodiac: str | None
    source: str


def load_targets() -> SheetTargets | None:
    if not SHEET_FILE.exists():
        return None
    try:
        raw = json.loads(SHEET_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(raw, dict):
        return None
    tail = _clean_tail(raw.get("tail"))
    zodiac = _clean_zodiac(raw.get("zodiac"))
    source = raw.get("source") if isinstance(raw.get("source"), str) else ""
    if tail is None and zodiac is None:
        return None
    return SheetTargets(tail=tail, zodiac=zodiac, source=source)


def save_targets(targets: SheetTargets) -> None:
    APP_DIR.mkdir(mode=0o700, parents=True, exist_ok=True)
    payload = {"tail": targets.tail, "zodiac": targets.zodiac, "source": targets.source}
    temp = SHEET_FILE.with_suffix(".tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    _restrict(temp)
    temp.replace(SHEET_FILE)
    _restrict(SHEET_FILE)


def import_targets(path: Path) -> SheetTargets:
    """Current H1 tail and I1 zodiac. Formulas have no cached value until Excel saves, so calculate them."""
    try:
        workbook = load_workbook(path, data_only=False)
    except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ValueError("请选择 .xlsx 文件") from exc
    worksheet = workbook["记录"] if "记录" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    tail, zodiac = _h1_i1(worksheet)
    if tail is None and zodiac is None:
        raise ValueError("表里还没有特码，算不出 H1 和 I1。")
    return SheetTargets(tail=tail, zodiac=zodiac, source=path.name)


def zero_trigger_issue(path: Path, written_issues: tuple[str, ...],
                       current_issue: str | None = None) -> str | None:
    """Return the newly completed latest issue only when its logical D value is zero.

    openpyxl does not calculate formulas, so reproduce the D-column rule from the
    actual B values instead of trusting a possibly stale cached formula result.
    """
    wanted = {str(value) for value in written_issues}
    try:
        workbook = load_workbook(path, data_only=False, read_only=True)
    except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ValueError("自动投注预警值检查失败：Excel 无法读取") from exc
    try:
        worksheet = workbook["记录"] if "记录" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        zodiac_of = _zodiac_map(worksheet)
        tail_last: dict[str, int] = {}
        zodiac_last: dict[str, int] = {}
        latest: tuple[str, bool] | None = None
        position = 0
        for row in range(3, (worksheet.max_row or 3) + 1):
            issue = _issue_text(worksheet.cell(row, 1).value)
            number = worksheet.cell(row, 2).value
            if issue is None or isinstance(number, bool) or not isinstance(number, (int, float)):
                continue
            number = int(number)
            if not 1 <= number <= 49:
                continue
            position += 1
            tail, zodiac = str(number)[-1], zodiac_of.get(number)
            target_tail = min(tail_last, key=tail_last.get) if tail_last else None
            target_zodiac = min(zodiac_last, key=zodiac_last.get) if zodiac_last else None
            is_zero = tail == target_tail or zodiac is not None and zodiac == target_zodiac
            latest = (issue, is_zero)
            tail_last[tail] = position
            if zodiac is not None:
                zodiac_last[zodiac] = position
        if latest is None or wanted and latest[0] not in wanted:
            return None
        if current_issue is not None and latest[0] != str(current_issue):
            return None
        # With new writes, only the newest completed row may arm a campaign.
        # With no writes, this also restores an unconsumed latest-zero trigger
        # after an application restart; AutoBetRunner deduplicates its issue.
        return latest[0] if latest[1] else None
    finally:
        workbook.close()


def d_formula(row: int) -> str:
    """D resets to 0 when this row's tail or zodiac hits the target from the rows above.

    H1 and I1 already include the current row, so they cannot be used here.
    The target is the same rule as H1/I1, calculated only through the previous row.
    """
    if row < 4:
        return '=IF(B3="",1,1)'
    tail = _target_formula("F", TAILS, f"F{row - 1}")[1:]
    zodiac = _target_formula("G", ZODIACS, f"G{row - 1}")[1:]
    return (
        f'=IF(B{row}="",IF(ISNUMBER(D{row - 1}),D{row - 1}+1,1),'
        f'IF(OR(AND(F{row}<>"",F{row}={tail}),AND(G{row}<>"",G{row}={zodiac})),'
        f'0,IF(ISNUMBER(D{row - 1}),D{row - 1}+1,1)))'
    )


def apply_d_formulas(worksheet, first: int = 3, last: int = 498) -> None:
    for row in range(first, last + 1):
        worksheet.cell(row, 4).value = d_formula(row)


def write_template(path: Path) -> None:
    base = Path(__file__).resolve().parent / "assets" / "sheet-template.xlsx"
    if not base.is_file():
        raise FileNotFoundError("缺少与源表格式一致的模板文件")
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(base, path)


def _target_formula(column: str, tokens: list[str], end: str) -> str:
    terms = ",".join(
        f'IF(COUNTIF(${column}$3:{end},"{token}")=0,10^9,'
        f'LOOKUP(2,1/(${column}$3:{end}="{token}"),ROW(${column}$3:{end})))'
        for token in tokens
    )
    return f'=IFERROR(INDEX(${column}$3:{end},MIN({terms})-ROW(${column}$3)+1),"")'


def _h1_i1(worksheet) -> tuple[str | None, str | None]:
    """Same result as H1 and I1: first tail and zodiac that never appear again below."""
    zodiac_of = _zodiac_map(worksheet)
    tails: list[str] = []
    zodiacs: list[str] = []
    for row in range(3, (worksheet.max_row or 3) + 1):
        number = worksheet.cell(row, 2).value
        if isinstance(number, str) or isinstance(number, bool) or not isinstance(number, (int, float)):
            continue
        tails.append(str(int(number))[-1])
        zodiacs.append(zodiac_of.get(int(number), ""))
    tail = next((value for index, value in enumerate(tails) if value not in tails[index + 1 :]), None)
    zodiac = next(
        (
            value
            for index, value in enumerate(zodiacs)
            if value and value not in zodiacs[index + 1 :]
        ),
        None,
    )
    if tail is None and zodiac is None:
        return _clean_tail(_literal(worksheet["H1"].value)), _clean_zodiac(_literal(worksheet["I1"].value))
    return tail, zodiac


def _zodiac_map(worksheet) -> dict[int, str]:
    found: dict[int, str] = {}
    for row in range(3, (worksheet.max_row or 3) + 1):
        number = worksheet.cell(row, 17).value
        name = worksheet.cell(row, 18).value
        if isinstance(number, bool) or not isinstance(number, (int, float)):
            continue
        if isinstance(name, str) and name.strip() in ZODIACS:
            found[int(number)] = name.strip()
    return found


def _literal(value):
    text = getattr(value, "text", None)
    if isinstance(text, str):
        return None
    return value


def _issue_text(value) -> str | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str) and value.strip().isdigit():
        return value.strip()
    return None


def _clean_tail(value) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        text = str(int(value))
    else:
        text = str(value).strip()
    if text in HEADER_SKIP or text == "":
        return None
    digit = text[-1]
    if digit in TAILS:
        return digit
    return None


def _clean_zodiac(value) -> str | None:
    if not isinstance(value, str):
        return None
    text = value.strip()
    if text in ZODIACS:
        return text
    return None
