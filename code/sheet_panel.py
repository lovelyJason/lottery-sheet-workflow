from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QFileDialog, QFrame, QGridLayout, QHBoxLayout, QLabel,
    QPushButton, QSizePolicy, QTabWidget, QVBoxLayout, QWidget,
)

from history_panel import HistoryPanel
from history_excel import PLAY_TWO_HEADERS
from play_options import PLAY2_LOGIC_OPTIONS, PLAY_OPTIONS
from sheet_book import write_template
from theme import DRAW_CONFIG_STYLE
from ui_common import hug, set_badge, unread_mark, _notice

class SheetPanel(QWidget):
    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setObjectName("drawConfig")
        self.setStyleSheet(DRAW_CONFIG_STYLE)
        self.setAttribute(Qt.WA_StyledBackground, True)
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(14)
        self.tabs = QTabWidget()
        self.tabs.setObjectName("configTabs")
        self.history_panel = HistoryPanel(self)
        self.tabs.addTab(self.history_panel, "历史补录")
        details = QWidget()
        self.tabs.addTab(details, "表格信息与模板")
        outer.addWidget(self.tabs, 1)
        root = QVBoxLayout(details)
        root.setContentsMargins(22, 20, 22, 18)
        root.setSpacing(12)
        root.setAlignment(Qt.AlignTop)

        head = QHBoxLayout()
        head.setSpacing(8)
        self.unread = unread_mark()
        title = QLabel("表格信息")
        title.setObjectName("heading")
        self.badge = QLabel("")
        self.badge.setObjectName("badge")
        self.badge.setProperty("state", "ready")
        self.badge.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        self.badge.setAlignment(Qt.AlignCenter)
        self.badge.hide()
        head.addWidget(self.unread, alignment=Qt.AlignVCenter)
        head.addWidget(title)
        head.addStretch(1)
        head.addWidget(self.badge)
        root.addLayout(head)

        self.hint = QLabel()
        self.hint.setObjectName("hint")
        self.hint.setWordWrap(True)
        root.addWidget(self.hint)

        self.box = QWidget()
        grid = QGridLayout(self.box)
        grid.setContentsMargins(0, 4, 0, 4)
        grid.setHorizontalSpacing(16)
        grid.setVerticalSpacing(8)
        grid.setColumnStretch(1, 1)
        self.tail_caption, self.tail = self._value(grid, 0, "H1 尾数")
        self.zodiac_caption, self.zodiac = self._value(grid, 1, "I1 生肖")
        self.logic_caption, self.logic = self._value(grid, 2, "当前逻辑")
        self.source_caption, self.source = self._value(grid, 3, "来源")
        root.addWidget(self.box)

        actions = QHBoxLayout()
        actions.setSpacing(10)
        template_btn = hug(QPushButton("下载 Excel 模板"))
        template_btn.setToolTip("按首页当前选择下载玩法一或玩法二模板")
        template_btn.clicked.connect(self._download)
        self.template_btn = template_btn
        actions.addWidget(template_btn)
        actions.addStretch(1)
        root.addLayout(actions)
        self.history_panel.state_changed.connect(self._sync_controls)
        self.history_panel.workbook_changed.connect(self.refresh)
        self.refresh()

    def _sync_controls(self, active: bool) -> None:
        self.template_btn.setEnabled(not active)

    def showEvent(self, event) -> None:
        self.refresh()
        self.history_panel._controls()
        super().showEvent(event)

    def _value(self, grid: QGridLayout, row: int, label: str) -> tuple[QLabel, QLabel]:
        caption = QLabel(label)
        caption.setObjectName("caption")
        caption.setMinimumWidth(88)
        value = QLabel("—")
        value.setObjectName("value")
        value.setTextFormat(Qt.PlainText)
        value.setWordWrap(True)
        value.setTextInteractionFlags(Qt.TextSelectableByMouse)
        grid.addWidget(caption, row, 0)
        grid.addWidget(value, row, 1)
        return caption, value

    def refresh(self) -> None:
        path = self.history_panel.config.get("workbook", "")
        ready = False
        hint = "请先在“历史补录”选择 Excel，这里会自动显示同一份表格的信息。"
        play_mode = self.history_panel._play_mode
        logic_key = self.history_panel._play2_logic
        self._set_info_layout(play_mode, logic_key)
        if path:
            try:
                workbook = load_workbook(path, data_only=True, read_only=True)
                try:
                    sheet = workbook["记录"] if "记录" in workbook.sheetnames else workbook.worksheets[0]
                    actual_kind = self._template_kind(sheet)
                    tail, zodiac = sheet["H1"].value, sheet["I1"].value
                finally:
                    workbook.close()
                expected_kind = "play_two" if play_mode == "play2" else "play_one"
                matches = actual_kind == expected_kind
                if actual_kind == "play_two":
                    self.tail.setText("玩法二")
                    self.zodiac.setText("B:H（平1至平6、特码）")
                    logic = next(
                        item.label for item in PLAY2_LOGIC_OPTIONS if item.key == logic_key
                    )
                    self.logic.setText(logic)
                    hint = (
                        "当前模板按 A 列期号匹配，B:H 仅补录空白开奖数据；"
                        "J/K 预警和 O/P/R 投注目标由程序按所选逻辑复算。"
                    )
                else:
                    self.tail.setText(str(tail) if tail is not None else "—")
                    self.zodiac.setText(str(zodiac) if zodiac is not None else "—")
                    hint = "读取当前补录文件的 H1 / I1；公式无缓存时显示 —，请用 Excel 计算并保存。"
                self.source.setText(Path(path).name)
                self.source.setToolTip(path)
                ready = matches
                if not matches:
                    selected = next(item.label for item in PLAY_OPTIONS if item.key == play_mode)
                    actual = "玩法二" if actual_kind == "play_two" else "玩法一"
                    hint = f"当前选择{selected}，但绑定的是{actual}模板；请重新选择对应 Excel。"
            except Exception:
                hint = "所选 Excel 暂时不可读，请检查文件或在“历史补录”重新选择。"
        self.unread.setVisible(not ready)
        set_badge(self.badge, ready, "已绑定")
        self.hint.setText(hint)
        self.box.setVisible(ready)

    def _set_info_layout(self, play_mode: str, logic_key: str) -> None:
        is_play2 = play_mode == "play2"
        self.tail_caption.setText("模板" if is_play2 else "H1 尾数")
        self.zodiac_caption.setText("补录列" if is_play2 else "I1 生肖")
        self.logic_caption.setVisible(is_play2)
        self.logic.setVisible(is_play2)
        if is_play2:
            logic = next(item.label for item in PLAY2_LOGIC_OPTIONS if item.key == logic_key)
            self.logic.setText(logic)

    @staticmethod
    def _template_kind(sheet) -> str:
        for row in range(1, min(sheet.max_row, 20) + 1):
            headers = tuple(sheet.cell(row, column).value for column in range(1, 9))
            if headers == PLAY_TWO_HEADERS:
                return "play_two"
        return "play_one"

    def _download(self) -> None:
        downloads = Path.home() / "Downloads"
        folder = downloads if downloads.is_dir() else Path.home()
        path, _ = QFileDialog.getSaveFileName(
            self,
            "保存 Excel 模板",
            str(folder / (
                "玩法二生肖投注模板.xlsx"
                if self.history_panel._play_mode == "play2"
                else "玩法一投注记录模板.xlsx"
            )),
            "Excel (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        try:
            write_template(Path(path), self.history_panel._play_mode)
        except (OSError, ValueError) as exc:
            _notice(self, "保存失败", str(exc) or "模板没有写入", warning=True)
            return
        if self.history_panel._play_mode == "play2":
            message = (
                "玩法二模板已保存。A 列是期数，程序会把平1至平6和特码"
                "补录到 B:H；启动前可选择逻辑一或逻辑二。"
            )
        else:
            message = (
                "玩法一模板已保存。只需填写期数和特码。新开奖的尾数对上当前 H，"
                "或生肖对上当前 I 时，这一期的 D 列会自动变成 0。"
            )
        _notice(
            self,
            "模板已保存",
            message,
        )
