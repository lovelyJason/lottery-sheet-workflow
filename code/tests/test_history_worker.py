import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication
from bet_client import BetError, BetOutcome, BetPlan
from history_worker import HistoryWorker

class WorkerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QCoreApplication.instance() or QCoreApplication([])
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/"book_已完成.xlsx"
        w=Workbook(); s=w.active; s.append(["期数","特码"])
        s.append(["115054996",None]); s.append(["115054980",None]); s.append(["115054997",None])
        w.save(self.path); w.close()
        self.worker=HistoryWorker("https://web.example.test",{},str(self.path),"2026-09-28")
        self.results=[];self.errors=[]
        self.worker.succeeded.connect(self.results.append)
        self.worker.failed.connect(self.errors.append)
        settings = patch("history_worker.load_settings", return_value=SimpleNamespace(
            profit_halt_date="", profit_halt_reason=""
        ))
        settings.start()
        self.addCleanup(settings.stop)
    def test_fetch_older_only_when_pending(self):
        pages=[[{"issue":"115054996","special_code":33}],[{"issue":"115054980","special_code":24}]]
        with patch("history_worker.HistoryClient") as client:
            client.return_value.fetch_page.side_effect=pages
            self.worker.run()
            self.assertEqual(client.return_value.fetch_page.call_count,2)
        self.assertFalse(self.errors)
        self.assertEqual(self.results[0]["report"].written,2)
        w=load_workbook(self.path)
        self.assertIsNone(w.active["B4"].value); w.close()
    def test_page_one_sufficient(self):
        with patch("history_worker.HistoryClient") as client:
            client.return_value.fetch_page.return_value=[
                {"issue":"115054996","special_code":33},{"issue":"115054980","special_code":24}]
            self.worker.run()
            self.assertEqual(client.return_value.fetch_page.call_count,1)

    def test_repeated_page_cancels_no_partial_write(self):
        before=self.path.read_bytes()
        with patch("history_worker.HistoryClient") as client:
            client.return_value.fetch_page.return_value=[{"issue":"115054996","special_code":33}]
            self.worker.run()
        self.assertTrue(self.errors);self.assertFalse(self.results)
        self.assertEqual(self.path.read_bytes(),before)
    def test_cancel_no_write(self):
        before=self.path.read_bytes()
        self.worker.cancel.set()
        with patch("history_worker.HistoryClient") as client:
            self.worker.run()
            client.return_value.fetch_page.assert_not_called()
        self.assertEqual(self.path.read_bytes(),before)

    def test_auto_bet_result_has_its_own_log_signal(self):
        worker = HistoryWorker("https://web.example.test", {}, str(self.path), "2026-09-28",
                               bet_plan=BetPlan(2, 3, "2", "猴", 4))
        bet_logs = []
        worker.bet_succeeded.connect(bet_logs.append)
        with patch("history_worker.HistoryClient") as history, \
             patch("history_worker.evaluate_bet_rule") as rule, \
             patch("history_worker.AutoBetRunner") as runner:
            history.return_value.fetch_page.return_value = [
                {"issue": "115054996", "special_code": 33},
                {"issue": "115054980", "special_code": 24},
            ]
            rule.return_value = SimpleNamespace(
                tail="2", zodiac="猴", trigger_issue="115055011",
                rule_key="play1", label="玩法一", exclusion_text="2尾/猴",
            )
            runner.return_value.run_once.return_value = BetOutcome(
                "115055012", 3, 25, 1, 2, "投注成功", 0,
                (("1", 2), ("5", 3), ("36", 20))
            )
            worker.run()
        self.assertEqual(len(bet_logs), 1)
        self.assertIn("特码B", bet_logs[0])
        self.assertIn("115055012期，投注了1、5、36，共3个号码", bet_logs[0])
        self.assertIn("每个号码积分各是2、3、20", bet_logs[0])
        self.assertIn("本期合计25积分", bet_logs[0])
        submitted_plan = runner.return_value.run_once.call_args.args[0]
        self.assertEqual(submitted_plan.start_offset, 4)
        self.assertEqual(
            runner.return_value.run_once.call_args.kwargs["latest_result_issue"],
            "115054996",
        )

    def test_auto_bet_business_error_does_not_erase_excel_success(self):
        worker = HistoryWorker("https://web.example.test", {}, str(self.path), "2026-09-28",
                               bet_plan=BetPlan(1, 1, "2", "猴"))
        successes, bet_errors, history_errors = [], [], []
        worker.succeeded.connect(successes.append)
        worker.bet_failed.connect(bet_errors.append)
        worker.failed.connect(history_errors.append)
        with patch("history_worker.HistoryClient") as history, \
             patch("history_worker.evaluate_bet_rule") as rule, \
             patch("history_worker.AutoBetRunner") as runner:
            history.return_value.fetch_page.return_value = [
                {"issue": "115054996", "special_code": 33},
                {"issue": "115054980", "special_code": 24},
            ]
            rule.return_value = SimpleNamespace(
                tail="2", zodiac="猴", trigger_issue="115055011",
                rule_key="play1", label="玩法一", exclusion_text="2尾/猴",
            )
            runner.return_value.run_once.side_effect = BetError("积分不足,投注失败")
            worker.run()
        self.assertTrue(successes)
        self.assertFalse(history_errors)
        self.assertEqual(bet_errors, ["自动投注失败：积分不足,投注失败"])

    def test_empty_daily_sheet_waits_after_sync_instead_of_failing_bet(self):
        worker = HistoryWorker(
            "https://web.example.test", {}, str(self.path), "2026-10-01",
            bet_plan=BetPlan(3, 40, None, None, 1, (40, 40, 40)),
        )
        successes, progress, bet_errors = [], [], []
        worker.succeeded.connect(successes.append)
        worker.progress.connect(progress.append)
        worker.bet_failed.connect(bet_errors.append)
        with patch("history_worker.HistoryClient") as history, \
                patch("history_worker.AutoBetRunner") as runner:
            # A valid result exists, but it precedes the first issue in this
            # freshly cleared daily sheet, so no B cell can be filled yet.
            history.return_value.fetch_page.side_effect = [[
                {"issue": "115055419", "special_code": 16},
            ], []]
            worker.run()

        self.assertTrue(successes)
        self.assertFalse(bet_errors)
        self.assertIn(
            "表格暂无已补录特码，等待开奖结果后再判断自动投注", progress
        )
        runner.return_value.run_once.assert_not_called()

    def test_profit_halt_is_rechecked_before_betting(self):
        worker = HistoryWorker("https://web.example.test", {}, str(self.path), "2026-09-28",
                               bet_plan=BetPlan(1, 1, "2", "猴"))
        successes, progress = [], []
        worker.succeeded.connect(successes.append)
        worker.progress.connect(progress.append)
        with patch("history_worker.HistoryClient") as history, \
             patch("history_worker.AutoBetRunner") as runner, \
             patch("history_worker.site_day", return_value="2026-09-29"), \
             patch("history_worker.load_settings", return_value=SimpleNamespace(
                 profit_halt_date="2026-09-29", profit_halt_reason="今日亏损达到停止值"
             )):
            history.return_value.fetch_page.return_value = [
                {"issue": "115054996", "special_code": 33},
                {"issue": "115054980", "special_code": 24},
            ]
            worker.run()
        runner.assert_not_called()
        self.assertTrue(successes)
        self.assertIn("今日亏损达到停止值", progress)

    def test_profit_is_automatically_checked_and_halts_before_betting(self):
        worker = HistoryWorker("https://web.example.test", {}, str(self.path), "2026-09-30",
                               bet_plan=BetPlan(1, 2, "2", "猴"))
        checked, successes, errors = [], [], []
        worker.profit_checked.connect(
            lambda value, halted, reason: checked.append((value, halted, reason))
        )
        worker.succeeded.connect(successes.append)
        worker.profit_failed.connect(errors.append)
        settings = SimpleNamespace(
            profit_halt_date="", profit_halt_reason="",
            profit_limit=6000, loss_limit=6000,
        )
        with patch("history_worker.HistoryClient") as history, \
             patch("history_worker.AutoBetRunner") as runner, \
             patch("history_worker.load_settings", return_value=settings), \
             patch("history_worker.save_profit_snapshot", return_value=(
                 True, "今日亏损 8552.6 已达到停止值 6000"
             )) as save_snapshot:
            history.return_value.fetch_page.return_value = [
                {"issue": "115054996", "special_code": 33},
                {"issue": "115054980", "special_code": 24},
            ]
            runner.return_value.client.today_profit.return_value = -8552.6
            worker.run()
        save_snapshot.assert_called_once_with(-8552.6)
        runner.return_value.run_once.assert_not_called()
        self.assertEqual(checked, [(
            -8552.6, True, "今日亏损 8552.6 已达到停止值 6000"
        )])
        self.assertTrue(successes)
        self.assertFalse(errors)

    def test_profit_check_failure_skips_bet_for_that_cycle(self):
        worker = HistoryWorker("https://web.example.test", {}, str(self.path), "2026-09-30",
                               bet_plan=BetPlan(1, 2, "2", "猴"))
        failures = []
        worker.profit_failed.connect(failures.append)
        settings = SimpleNamespace(
            profit_halt_date="", profit_halt_reason="",
            profit_limit=6000, loss_limit=6000,
        )
        with patch("history_worker.HistoryClient") as history, \
             patch("history_worker.AutoBetRunner") as runner, \
             patch("history_worker.load_settings", return_value=settings):
            history.return_value.fetch_page.return_value = [
                {"issue": "115054996", "special_code": 33},
                {"issue": "115054980", "special_code": 24},
            ]
            runner.return_value.client.today_profit.side_effect = BetError("登录态失效")
            worker.run()
        runner.return_value.run_once.assert_not_called()
        self.assertEqual(
            failures, ["自动盈亏检查失败，本期停止投注：登录态失效"]
        )

if __name__ == "__main__":
    unittest.main()
