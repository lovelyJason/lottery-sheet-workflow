import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from sheet_book import d_formula, zero_trigger_issue


class ZeroTriggerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "book.xlsx"

    def _book(self, numbers):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["标题"])
        sheet.append(["期数", "特码"])
        for index, number in enumerate(numbers, 1):
            sheet.append([f"115055{index:03d}", number])
        # Minimal number→zodiac lookup used by the real workbook.
        for row, (number, zodiac) in enumerate(((1, "马"), (2, "蛇"), (3, "龙")), 3):
            sheet.cell(row, 17).value = number
            sheet.cell(row, 18).value = zodiac
        workbook.save(self.path)
        workbook.close()

    def test_new_latest_zero_arms_campaign(self):
        # Before issue 004, tail 1 is least recently seen; number 11 hits that tail.
        self._book([1, 2, 3, 11])
        self.assertEqual(zero_trigger_issue(self.path, ("115055004",)), "115055004")

    def test_latest_zero_can_restore_trigger_after_restart(self):
        self._book([1, 2, 3, 11])
        self.assertEqual(zero_trigger_issue(self.path, (), "115055004"), "115055004")

    def test_stale_workbook_zero_does_not_arm_current_site_issue(self):
        self._book([1, 2, 3, 11])
        self.assertIsNone(zero_trigger_issue(self.path, (), "115055099"))

    def test_new_latest_nonzero_does_not_arm(self):
        self._book([1, 2, 3, 14])
        self.assertIsNone(zero_trigger_issue(self.path, ("115055004",)))

    def test_old_gap_zero_does_not_arm_current_campaign(self):
        self._book([1, 2, 3, 11, 14])
        self.assertIsNone(zero_trigger_issue(self.path, ("115055004",)))

    def test_missing_draw_formula_keeps_d_counter_running(self):
        self.assertEqual(d_formula(3), '=IF(B3="",1,1)')
        self.assertTrue(d_formula(4).startswith(
            '=IF(B4="",IF(ISNUMBER(D3),D3+1,1),'
        ))


if __name__ == "__main__":
    unittest.main()
