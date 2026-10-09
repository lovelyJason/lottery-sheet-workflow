import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from bet_rules import (
    PLAY1, PLAY2_LOGIC1, PLAY2_LOGIC2, RULES,
    evaluate_bet_rule, rule_key,
)


class BetRuleAdapterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "play2.xlsx"

    def _book(self, rows, mapping):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["数据"])
        sheet.append(["期数", "平1", "平2", "平3", "平4", "平5", "平6",
                      "特码", "生肖", "尾/肖", "杀1肖"])
        for issue, regular, special in rows:
            sheet.append([issue, *regular, special])
        for index, (number, zodiac) in enumerate(mapping.items(), 3):
            sheet.cell(index, 24).value = number
            sheet.cell(index, 25).value = zodiac
        # Deliberately misleading formulas/cells: adapters must calculate from
        # literal draw data because openpyxl cannot refresh formula caches.
        sheet["O1"] = '=IF(TRUE,"9","")'
        sheet["P1"] = '=IF(TRUE,"猪","")'
        sheet["R1"] = '=IF(TRUE,"猪","")'
        workbook.save(self.path)
        workbook.close()

    def test_two_level_selection_maps_to_extensible_registry(self):
        self.assertEqual(rule_key("play1", "logic2"), PLAY1)
        self.assertEqual(rule_key("play2", "logic1"), PLAY2_LOGIC1)
        self.assertEqual(rule_key("play2", "logic2"), PLAY2_LOGIC2)
        self.assertEqual(set(RULES), {PLAY1, PLAY2_LOGIC1, PLAY2_LOGIC2})
        with self.assertRaisesRegex(ValueError, "玩法"):
            rule_key("future-play", "logic1")

    def test_play2_logic1_recomputes_j_and_current_o_p_targets(self):
        regular = (4, 5, 6, 7, 8, 9)
        self._book([
            ("115055001", regular, 1),
            ("115055002", regular, 2),
            ("115055003", regular, 3),
            ("115055004", regular, 11),
        ], {1: "马", 2: "蛇", 3: "龙", 11: "马"})

        decision = evaluate_bet_rule(
            PLAY2_LOGIC1, self.path, ("115055004",), "115055004"
        )

        self.assertEqual(decision.trigger_issue, "115055004")
        self.assertEqual((decision.tail, decision.zodiac), ("2", "蛇"))
        self.assertEqual(decision.exclusion_text, "2尾/蛇")

    def test_play2_logic1_ignores_an_old_zero(self):
        regular = (4, 5, 6, 7, 8, 9)
        self._book([
            ("115055001", regular, 1),
            ("115055002", regular, 2),
            ("115055003", regular, 3),
            ("115055004", regular, 11),
            ("115055005", regular, 14),
        ], {1: "马", 2: "蛇", 3: "龙", 11: "马", 14: "兔"})

        decision = evaluate_bet_rule(
            PLAY2_LOGIC1, self.path, ("115055004",), "115055005"
        )

        self.assertIsNone(decision.trigger_issue)

    def test_play2_logic2_recomputes_k_and_excludes_current_r_zodiac(self):
        # First row predicts total 30 -> 猴. The second special is also 猴,
        # therefore K reaches zero. Its own R prediction is total 37 -> 鸡,
        # which becomes the zodiac excluded by the triggered campaign.
        self._book([
            ("115055001", (1, 2, 3, 4, 5, 6), 1),
            ("115055002", (7, 8, 9, 10, 11, 12), 3),
        ], {1: "鼠", 3: "猴", 30: "猴", 37: "鸡"})

        decision = evaluate_bet_rule(
            PLAY2_LOGIC2, self.path, ("115055002",), "115055002"
        )

        self.assertEqual(decision.trigger_issue, "115055002")
        self.assertIsNone(decision.tail)
        self.assertEqual(decision.zodiac, "鸡")
        self.assertEqual(decision.exclusion_text, "鸡")

    def test_play2_rejects_a_play1_workbook_before_betting(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["期数", "特码"])
        sheet.append(["115055001", 3])
        workbook.save(self.path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "不是玩法二模板"):
            evaluate_bet_rule(PLAY2_LOGIC1, self.path, (), "115055001")


if __name__ == "__main__":
    unittest.main()
