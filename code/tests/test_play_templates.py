"""Packaged play templates and play/layout contract regressions."""
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from history_excel import PLAY_TWO_HEADERS, WorkbookSync
from history_panel import HistoryPanel
from play_options import template_kind
from sheet_book import write_template


class PlayTemplateTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)

    def test_play_two_download_preserves_headers_formulas_and_styles(self):
        output = self.root / "play2.xlsx"
        write_template(output, "play2")

        workbook = load_workbook(output, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            self.assertEqual(
                tuple(sheet.cell(2, column).value for column in range(1, 9)),
                PLAY_TWO_HEADERS,
            )
            self.assertEqual(sheet["K4"].value, "=IF(I4=R3,0,K3+1)")
            self.assertTrue(str(sheet["R4"].value).startswith("=VLOOKUP(Q4"))
            self.assertTrue(sheet["A2"].font.bold)
            self.assertEqual(sheet["A2"].alignment.horizontal, "center")
            self.assertEqual(sheet["A2"].border.left.style, "thin")
        finally:
            workbook.close()

        sync = WorkbookSync(output)
        try:
            self.assertEqual(sync.template_kind, "play_two")
        finally:
            sync.close()

    def test_play_one_download_remains_the_default(self):
        output = self.root / "play1.xlsx"
        write_template(output)
        sync = WorkbookSync(output)
        try:
            self.assertEqual(sync.template_kind, "play_one")
        finally:
            sync.close()

    def test_play_selection_requires_the_matching_template(self):
        self.assertEqual(template_kind("play1"), "play_one")
        self.assertEqual(template_kind("play2"), "play_two")
        HistoryPanel._validate_template("play_one", "play1")
        HistoryPanel._validate_template("play_two", "play2")
        with self.assertRaisesRegex(ValueError, "Excel 是玩法一模板"):
            HistoryPanel._validate_template("play_one", "play2")
        with self.assertRaisesRegex(ValueError, "Excel 是玩法二模板"):
            HistoryPanel._validate_template("play_two", "play1")

    def test_unknown_play_does_not_silently_download_the_wrong_template(self):
        with self.assertRaisesRegex(ValueError, "玩法无效"):
            write_template(self.root / "wrong.xlsx", "play3")


if __name__ == "__main__":
    unittest.main()
