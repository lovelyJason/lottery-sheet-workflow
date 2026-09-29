import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from history_excel import WorkbookSync, WorkbookError, import_source

def record(issue, code=33):
    return {"issue":str(issue),"special_code":code}

class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/"source.xlsx"
    def make(self, rows):
        w=Workbook();s=w.active;s.title="记录"
        s.append(["期数","特码","公式"])
        for row in rows:s.append(row)
        s["C2"]="=B2+1"
        s["B2"].fill=PatternFill("solid",fgColor="FFCC00")
        w.save(self.path);w.close()
        return self.path
    def read(self,path,col="B"):
        w=load_workbook(path,data_only=False)
        values=[w.active[f"{col}{r}"].value for r in range(2,w.active.max_row+1)]
        w.close()
        return values
    def sync(self,path):
        s=WorkbookSync(path);self.addCleanup(s.close)
        return s
    def test_full_and_three_digit_leading_zeros_preserve_original(self):
        self.make([["115055000",None],["999",None],["001",None]])
        raw=self.path.read_bytes();out=import_source(self.path)
        s=self.sync(out)
        self.assertFalse(s.needs_more([record(115055001,2),record(115055000,3),record(115054999,18)]))
        self.assertEqual(s.apply([record(115055001,2),record(115055000,3),record(115054999,18)]).written,3)
        self.assertEqual(self.read(out),[3,18,2]);self.assertEqual(self.path.read_bytes(),raw)
        self.assertEqual(self.sync(out).apply([record(115055001,2),record(115055000,3),record(115054999,18)]).written,0)
    def test_existing_zero_formula_and_text_are_not_empty(self):
        self.make([[115054996,0],[115054995,'=IF(1,"",1)'],[115054994,"已有"],[115054993," "]])
        s=self.sync(self.path)
        self.assertEqual(s.apply([record(i) for i in range(115054993,115054997)]).written,1)
        self.assertEqual(self.read(self.path),[0,'=IF(1,"",1)',"已有",33])
    def test_full_never_matches_suffix(self):
        self.make([["115053996",None]])
        s=self.sync(self.path);self.assertEqual(s.apply([record(115054996)]).written,0)
        self.assertTrue(s.needs_more([record(115054996)]))
    def test_future_no_paging_but_older_pages(self):
        self.make([["997",None]])
        s=self.sync(self.path);self.assertFalse(s.needs_more([record(115054996)]))
        self.make([["980",None]])
        s=self.sync(self.path);self.assertTrue(s.needs_more([record(115054996)]))
    def test_multiple_cycles_ambiguous(self):
        self.make([["996",None]])
        report=self.sync(self.path).apply([record(115054996),record(115053996,12)])
        self.assertEqual((report.written,report.ambiguous),(0,1))
    def test_full_suffix_duplicate_targets(self):
        self.make([["996",None],[115054996,None]])
        report=self.sync(self.path).apply([record(115054996)])
        self.assertEqual((report.written,report.ambiguous),(0,2))
    def test_conflicting_api_data_not_written(self):
        self.make([[115054996,None]])
        report=self.sync(self.path).apply([record(115054996),record(115054996,12)])
        self.assertEqual((report.written,report.ambiguous),(0,1))
    def test_styles_and_formulas_preserved(self):
        self.make([[115054996,None]])
        s=self.sync(self.path);s.apply([record(115054996)])
        w=load_workbook(self.path);self.assertEqual(w.active["C2"].value,"=B2+1")
        self.assertEqual(w.active["B2"].fill.fgColor.rgb,"00FFCC00");w.close()
    def test_legacy_d_formula_is_upgraded_so_missing_draw_increments(self):
        w=Workbook();s=w.active;s.title="记录"
        s.append(["标题"]);s.append(["期数","特码",None,"预警"])
        s.append(["115054996",None,None,'=IF(B3="","",1)'])
        s.append(["115054997",None,None,'=IF(B4="","",IF(TRUE,0,D3+1))'])
        w.save(self.path);w.close()
        self.sync(self.path).apply([])
        w=load_workbook(self.path,data_only=False)
        self.assertEqual(w["记录"]["D3"].value,'=IF(B3="",1,1)')
        self.assertEqual(
            w["记录"]["D4"].value,
            '=IF(B4="",IF(ISNUMBER(D3),D3+1,1),IF(TRUE,0,D3+1))',
        )
        w.close()
    def test_disk_change_prevents_overwrite(self):
        self.make([[115054996,None]])
        s=self.sync(self.path)
        w=load_workbook(self.path);w.active["B2"]=49;w.save(self.path);w.close()
        with self.assertRaises(WorkbookError):s.apply([record(115054996)])
        self.assertEqual(self.read(self.path),[49])
    def test_atomic_failure_keeps_original_and_cleans_temp(self):
        self.make([[115054996,None]]);raw=self.path.read_bytes();s=self.sync(self.path)
        with patch("history_excel.os.replace",side_effect=PermissionError):
            with self.assertRaises(WorkbookError):s.apply([record(115054996)])
        self.assertEqual(self.path.read_bytes(),raw)
        self.assertFalse(list(self.path.parent.glob(".history-*")))
    def test_excel_lock(self):
        self.make([[115054996,None]])
        self.path.with_name("~$"+self.path.name).touch()
        with self.assertRaises(WorkbookError):WorkbookSync(self.path)
    def test_no_change_does_not_rewrite(self):
        self.make([[115054996,33]]);raw=self.path.read_bytes()
        self.assertEqual(self.sync(self.path).apply([record(115054996)]).written,0)
        self.assertEqual(self.path.read_bytes(),raw)
    def test_existing_copy_not_overwritten(self):
        self.make([[115054996,None]]);out=import_source(self.path);before=out.read_bytes()
        with self.assertRaises(WorkbookError):import_source(self.path)
        self.assertEqual(out.read_bytes(),before)
    def test_formula_issue_not_guessed(self):
        self.make([['=115054996',None]])
        self.assertEqual(self.sync(self.path).pending_count,0)
    def test_empty_batch_and_invalid_number(self):
        self.make([[115054996,None]]);s=self.sync(self.path)
        self.assertFalse(s.needs_more([]));self.assertEqual(s.apply([]).written,0)
        for code in (0,50,True):
            with self.subTest(code=code),self.assertRaises(WorkbookError):s.apply([record(115054996,code)])

if __name__=="__main__":unittest.main()
