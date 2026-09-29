"""Single chooser and copy/direct mode regressions; isolated temporary files only."""
from contextlib import ExitStack
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QPushButton

import history_config
from history_excel import WorkbookSync
from history_panel import HistoryPanel
from sheet_panel import SheetPanel
from settings_store import Settings


class WorkbookSelectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)

    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.folder = Path(self.stack.enter_context(tempfile.TemporaryDirectory())).resolve()
        self.source = self.folder / "test.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "记录"
        ws.append(["期数", "特码"])
        ws.append(["996", None])
        ws.append(["995", 47])
        ws["H1"] = 2
        ws["I1"] = "猴"
        wb.save(self.source)
        wb.close()
        self.original = self.source.read_bytes()
        self.stack.enter_context(patch.object(history_config, "CONFIG", self.folder / "history.json"))
        self.stack.enter_context(patch("app_log.info"))
        self.stack.enter_context(patch("app_log.error"))
        self.dialog = SheetPanel()
        self.panel = self.dialog.history_panel
        self.addCleanup(self.dialog.deleteLater)
        self.addCleanup(self.dialog.close)

    def choose(self, path=None):
        with patch.object(self.panel, "_select", return_value=str(path or self.source)):
            self.panel.choose.click()

    def write_result(self):
        sync = WorkbookSync(Path(self.panel.config["workbook"]))
        try:
            return sync.apply([{"issue": "115054996", "special_code": 33},
                               {"issue": "115054995", "special_code": 1}])
        finally:
            sync.close()

    def test_single_chooser_and_no_details_import(self):
        buttons = [b.text() for b in self.dialog.findChildren(QPushButton)]
        self.assertEqual(buttons.count("选择 Excel"), 1)
        self.assertNotIn("续用工作副本", buttons)
        self.assertNotIn("导入 Excel", buttons)
        self.assertIn("下载 Excel 模板", buttons)
        self.assertTrue(self.panel.use_copy.isChecked())

    def test_copy_keeps_source_unchanged_and_fills_only_blank(self):
        self.choose()
        self.assertEqual(self.panel.config["workbook"], str(self.folder / "test_已完成.xlsx"))
        self.assertEqual(self.write_result().written, 1)
        self.assertEqual(self.source.read_bytes(), self.original)
        wb = load_workbook(self.panel.config["workbook"])
        self.assertEqual(wb["记录"]["B2"].value, 33)
        self.assertEqual(wb["记录"]["B3"].value, 47)
        wb.close()
        self.assertEqual(self.write_result().written, 0)

    def test_unchecked_writes_the_selected_source(self):
        self.panel.use_copy.setChecked(False)
        self.choose()
        self.assertEqual(self.panel.config["workbook"], str(self.source))
        self.assertEqual(self.write_result().written, 1)
        wb = load_workbook(self.source)
        self.assertEqual(wb["记录"]["B2"].value, 33)
        self.assertEqual(wb["记录"]["B3"].value, 47)
        wb.close()
        self.assertFalse((self.folder / "test_已完成.xlsx").exists())
        self.assertIn("直接修改", self.panel.file_hint.text())

    def test_toggle_and_restart_reuse_known_copy(self):
        self.choose()
        copied = self.panel.config["workbook"]
        self.write_result()
        copy_bytes = Path(copied).read_bytes()
        self.panel.use_copy.setChecked(False)
        self.assertEqual(self.panel.config["workbook"], str(self.source))
        restored = HistoryPanel()
        try:
            self.assertFalse(restored.use_copy.isChecked())
            restored.use_copy.setChecked(True)
            self.assertEqual(restored.config["workbook"], copied)
            self.assertEqual(Path(copied).read_bytes(), copy_bytes)
            self.assertEqual(history_config.load_config()["source"], str(self.source))
            self.assertTrue(history_config.load_config()["use_copy"])
        finally:
            restored.deleteLater()

    def test_unrelated_existing_copy_not_overwritten_and_mode_restored(self):
        other = self.folder / "test_已完成.xlsx"
        other.write_bytes(b"unrelated file")
        self.panel.use_copy.setChecked(False)
        self.choose()
        previous = dict(self.panel.config)
        self.panel.use_copy.setChecked(True)
        self.assertFalse(self.panel.use_copy.isChecked())
        self.assertEqual(self.panel.config, previous)
        self.assertEqual(other.read_bytes(), b"unrelated file")
        self.assertIn("未覆盖", self.panel.note.text())

    def test_copy_means_copy_even_when_selected_name_ends_completed(self):
        renamed = self.folder / "existing_已完成.xlsx"
        renamed.write_bytes(self.original)
        self.choose(renamed)
        self.assertNotEqual(self.panel.config["workbook"], str(renamed))
        self.write_result()
        self.assertEqual(renamed.read_bytes(), self.original)

    def test_cancel_and_invalid_selection_retain_binding(self):
        self.choose()
        before = dict(self.panel.config)
        with patch.object(self.panel, "_select", return_value=""):
            self.panel.choose.click()
        self.assertEqual(self.panel.config, before)
        invalid = self.folder / "invalid.xlsx"
        invalid.write_text("not excel")
        self.choose(invalid)
        self.assertEqual(self.panel.config, before)

    def test_save_failure_restores_binding_and_checkbox(self):
        self.choose()
        before = dict(self.panel.config)
        with patch("history_panel.save_config", side_effect=OSError("disk full")):
            self.panel.use_copy.setChecked(False)
        self.assertTrue(self.panel.use_copy.isChecked())
        self.assertEqual(self.panel.config, before)
        self.assertEqual(history_config.load_config()["workbook"], before["workbook"])

    def test_info_uses_same_workbook_and_refreshes_after_write_signal(self):
        self.choose()
        self.assertEqual(self.dialog.tail.text(), "2")
        self.assertEqual(self.dialog.zodiac.text(), "猴")
        self.assertEqual(self.dialog.source.text(), "test_已完成.xlsx")
        self.assertEqual(self.dialog.source.toolTip(), self.panel.config["workbook"])
        wb = load_workbook(self.panel.config["workbook"])
        wb["记录"]["H1"] = 7
        wb.save(self.panel.config["workbook"])
        wb.close()
        self.panel.workbook_changed.emit()
        self.assertEqual(self.dialog.tail.text(), "7")
        self.panel.use_copy.setChecked(False)
        self.assertEqual(self.dialog.tail.text(), "2")
        self.assertEqual(self.dialog.source.text(), "test.xlsx")

    def test_programmatic_mode_change_is_guarded_during_polling(self):
        self.choose()
        before = dict(self.panel.config)
        self.panel.running = True
        self.panel._controls()
        self.assertFalse(self.panel.use_copy.isEnabled())
        self.panel.use_copy.setChecked(False)
        self.assertTrue(self.panel.use_copy.isChecked())
        self.assertEqual(self.panel.config, before)
        self.assertTrue(self.panel.running)
        self.panel.running = False

    def test_direct_source_is_accepted_by_query_launcher(self):
        self.panel.use_copy.setChecked(False)
        self.choose()
        with patch("history_panel.load_settings", return_value=Settings(
                "https://web.example.test", None, None, 30)), \
                patch("history_panel.load", return_value={"token": "test-token"}), \
                patch("history_panel.HistoryWorker") as worker:
            self.panel._launch()
            worker.assert_called_once()
            self.assertEqual(worker.call_args.args[2], str(self.source))
            worker.return_value.start.assert_called_once()
            self.assertFalse(self.panel.use_copy.isEnabled())
            self.panel.worker = None
            self.panel._controls()

    def test_config_empty_paths_stay_empty_and_legacy_is_supported(self):
        history_config.save_config("", "", use_copy=False)
        self.assertEqual(history_config.load_config()["workbook"], "")
        self.assertFalse(history_config.load_config()["use_copy"])
        history_config.CONFIG.write_text('{"workbook": "legacy_已完成.xlsx"}')
        legacy = history_config.load_config()
        self.assertTrue(legacy["use_copy"])
        self.assertEqual(legacy["workbook"], "legacy_已完成.xlsx")


if __name__ == "__main__":
    unittest.main()
