"""Extensible workbook-driven betting rule adapters.

Formula cells are deliberately not used as inputs.  openpyxl preserves formulas
but cannot calculate them and may remove cached results after saving, so every
adapter reproduces the relevant workbook rule from literal draw cells.
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from sheet_book import NoSheetResults, ZODIACS, import_targets, zero_trigger_issue

PLAY1 = "play1"
PLAY2_LOGIC1 = "play2_logic1"
PLAY2_LOGIC2 = "play2_logic2"


@dataclass(frozen=True)
class BetRuleDecision:
    rule_key: str
    label: str
    trigger_issue: str | None
    tail: str | None
    zodiac: str | None

    @property
    def exclusion_text(self) -> str:
        if self.tail is not None:
            return f"{self.tail}尾/{self.zodiac or '—'}"
        return self.zodiac or "—"


class BetRuleAdapter(Protocol):
    key: str
    label: str

    def evaluate(self, path: Path, written_issues: tuple[str, ...],
                 current_issue: str | None) -> BetRuleDecision: ...


def rule_key(play_mode: object, play2_logic: object) -> str:
    """Map the persisted two-level GUI selection to one registry key."""
    if play_mode == "play1":
        return PLAY1
    if play_mode == "play2" and play2_logic == "logic1":
        return PLAY2_LOGIC1
    if play_mode == "play2" and play2_logic == "logic2":
        return PLAY2_LOGIC2
    raise ValueError("玩法或玩法二逻辑无效")


def evaluate_bet_rule(key: str, path: Path, written_issues: tuple[str, ...],
                      current_issue: str | None = None) -> BetRuleDecision:
    try:
        adapter = RULES[key]
    except KeyError:
        raise ValueError("自动投注规则无效") from None
    return adapter.evaluate(Path(path), written_issues, current_issue)


class Play1Rule:
    key = PLAY1
    label = "玩法一"

    def evaluate(self, path: Path, written_issues: tuple[str, ...],
                 current_issue: str | None) -> BetRuleDecision:
        targets = import_targets(path)
        trigger = zero_trigger_issue(path, written_issues, current_issue)
        return BetRuleDecision(
            self.key, self.label, trigger, targets.tail, targets.zodiac
        )


class _Play2Rule:
    key = ""
    label = ""

    def _load(self, path: Path):
        try:
            # Normal mode is intentional: the play-two formulas span hundreds
            # of rows and read-only random cell access is prohibitively slow.
            workbook = load_workbook(path, data_only=False, read_only=False)
        except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise ValueError("玩法二规则检查失败：Excel 无法读取") from exc
        worksheet = workbook["记录"] if "记录" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        expected = ("期数", "平1", "平2", "平3", "平4", "平5", "平6", "特码")
        actual = tuple(worksheet.cell(2, column).value for column in range(1, 9))
        if actual != expected:
            workbook.close()
            raise ValueError("当前 Excel 不是玩法二模板：需要 A-H 列为期数、平1至平6、特码")
        return workbook, worksheet


class Play2Logic1Rule(_Play2Rule):
    key = PLAY2_LOGIC1
    label = "玩法二 / 逻辑一"

    def evaluate(self, path: Path, written_issues: tuple[str, ...],
                 current_issue: str | None) -> BetRuleDecision:
        workbook, worksheet = self._load(path)
        try:
            zodiac_of = _zodiac_lookup(worksheet)
            rows = _completed_rows(worksheet, zodiac_of)
            if not rows:
                raise NoSheetResults("表里还没有完整的玩法二开奖结果。")
            tail_last: dict[str, int] = {}
            zodiac_last: dict[str, int] = {}
            latest_issue, latest_zero = "", False
            for position, row in enumerate(rows, 1):
                tail = str(row.special)[-1]
                target_tail = min(tail_last, key=tail_last.get) if tail_last else None
                target_zodiac = min(zodiac_last, key=zodiac_last.get) if zodiac_last else None
                latest_zero = tail == target_tail or row.special_zodiac == target_zodiac
                latest_issue = row.issue
                tail_last[tail] = position
                zodiac_last[row.special_zodiac] = position
            target_tail = min(tail_last, key=tail_last.get)
            target_zodiac = min(zodiac_last, key=zodiac_last.get)
            trigger = _trigger(latest_issue, latest_zero, written_issues, current_issue)
            return BetRuleDecision(
                self.key, self.label, trigger, target_tail, target_zodiac
            )
        finally:
            workbook.close()


class Play2Logic2Rule(_Play2Rule):
    key = PLAY2_LOGIC2
    label = "玩法二 / 逻辑二"

    def evaluate(self, path: Path, written_issues: tuple[str, ...],
                 current_issue: str | None) -> BetRuleDecision:
        workbook, worksheet = self._load(path)
        try:
            zodiac_of = _zodiac_lookup(worksheet)
            rows = _completed_rows(worksheet, zodiac_of)
            if not rows:
                raise NoSheetResults("表里还没有完整的玩法二开奖结果。")
            latest_issue, latest_zero = rows[0].issue, False
            previous_prediction = _prediction_zodiac(rows[0].regular, zodiac_of)
            latest_prediction = previous_prediction
            for row in rows[1:]:
                latest_zero = row.special_zodiac == previous_prediction
                latest_issue = row.issue
                latest_prediction = _prediction_zodiac(row.regular, zodiac_of)
                previous_prediction = latest_prediction
            trigger = _trigger(latest_issue, latest_zero, written_issues, current_issue)
            return BetRuleDecision(
                self.key, self.label, trigger, None, latest_prediction
            )
        finally:
            workbook.close()


@dataclass(frozen=True)
class _Play2Row:
    issue: str
    regular: tuple[int, ...]
    special: int
    special_zodiac: str


def _completed_rows(worksheet, zodiac_of: dict[int, str]) -> list[_Play2Row]:
    rows: list[_Play2Row] = []
    started = False
    for index in range(3, (worksheet.max_row or 3) + 1):
        issue = _issue_text(worksheet.cell(index, 1).value)
        values = tuple(_draw_number(worksheet.cell(index, column).value)
                       for column in range(2, 9))
        complete = issue is not None and all(value is not None for value in values)
        if not complete:
            if started:
                break
            continue
        started = True
        regular = tuple(value for value in values[:6] if value is not None)
        special = values[6]
        zodiac = zodiac_of.get(special) if special is not None else None
        if zodiac is None:
            raise ValueError(f"玩法二模板缺少特码 {special} 的生肖对应关系")
        rows.append(_Play2Row(issue, regular, special, zodiac))
    return rows


def _prediction_zodiac(regular: tuple[int, ...], zodiac_of: dict[int, str]) -> str:
    ordered = sorted(regular)
    fourth, fifth, sixth = ordered[3:6]
    value = ((fifth // 10 + fifth % 10) % 10 + fourth // 10
             + sixth // 10 + sixth % 10 + fourth + sixth + 9)
    zodiac = zodiac_of.get(value)
    if zodiac is None:
        raise ValueError(f"玩法二模板缺少合计 {value} 的生肖对应关系")
    return zodiac


def _zodiac_lookup(worksheet) -> dict[int, str]:
    result: dict[int, str] = {}
    for row in range(3, (worksheet.max_row or 3) + 1):
        number = _positive_integer(worksheet.cell(row, 24).value)
        zodiac = worksheet.cell(row, 25).value
        if number is not None and isinstance(zodiac, str) and zodiac.strip() in ZODIACS:
            result[number] = zodiac.strip()
    if not result:
        raise ValueError("玩法二模板缺少 X/Y 列生肖对应表")
    return result


def _trigger(issue: str, is_zero: bool, written_issues: tuple[str, ...],
             current_issue: str | None) -> str | None:
    if not is_zero:
        return None
    written = {str(value) for value in written_issues}
    if written and issue not in written:
        return None
    if current_issue is not None and issue != str(current_issue):
        return None
    return issue


def _draw_number(value) -> int | None:
    number = _positive_integer(value)
    return number if number is not None and number <= 49 else None


def _positive_integer(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value)
    return None


def _issue_text(value) -> str | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int) and value >= 0:
        return str(value)
    if isinstance(value, float) and value.is_integer() and value >= 0:
        return str(int(value))
    if isinstance(value, str) and value.strip().isdigit():
        return value.strip()
    return None


RULES: dict[str, BetRuleAdapter] = {
    adapter.key: adapter
    for adapter in (Play1Rule(), Play2Logic1Rule(), Play2Logic2Rule())
}
